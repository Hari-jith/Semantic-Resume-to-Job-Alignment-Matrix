"""
DOCX text extraction using python-docx.

Extracts both regular paragraph text and text inside tables, since some
resume templates use tables for layout (e.g. skills grids, two-column
experience sections) and that text would otherwise be silently dropped.
"""

import io

from docx import Document

from app.parsers.file_validator import FileValidationError

MIN_MEANINGFUL_CHARS = 40


def extract_text_from_docx(file_bytes: bytes) -> str:
    """
    Extracts and returns all text from a DOCX file.
    Raises FileValidationError with a user-facing message on failure.
    """
    try:
        document = Document(io.BytesIO(file_bytes))
    except Exception as exc:
        # python-docx raises various exceptions (BadZipFile, KeyError, etc.)
        # for corrupted files or files that aren't actually valid DOCX.
        raise FileValidationError(
            "Could not read the DOCX file. It may be corrupted or not a valid Word document."
        ) from exc

    text_parts = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text_parts.append(paragraph.text)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text_parts.append(cell.text)

    full_text = "\n".join(text_parts).strip()

    if len(full_text) < MIN_MEANINGFUL_CHARS:
        raise FileValidationError(
            "No readable text could be extracted from this DOCX file. "
            "The document may be empty or contain only images."
        )

    return full_text
